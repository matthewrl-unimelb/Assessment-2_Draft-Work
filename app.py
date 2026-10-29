"""
app.py - Bank Merger Regulatory Check (Streamlit).

Run locally:   streamlit run app.py
Structure:     Home + 7 steps, a tab bar at the top, Back / Next on every step.
Rule of thumb: engine.py calculates and decides; ai.py reads and explains; this file shows things.
"""

import hashlib
import io
import json
import os

import pandas as pd
import plotly.graph_objects as go
import streamlit as st

import ai
import engine as E
from library import Library

try:  # local development: read API keys from a .env file (never commit it)
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

st.set_page_config(page_title="Bank Merger Regulatory Check", page_icon="🏦", layout="wide")


# Streamlit reads "$...$" as maths, which garbles "$51bn ... $100bn". Escape dollar signs in all text
# we display (HTML tables and cards are left alone because markdown does not touch them).
def _escape_dollars(fn):
    import re

    def wrapper(body="", *args, **kwargs):
        if isinstance(body, str) and not body.lstrip().startswith(("<table", "<div", "<style")):
            body = re.sub(r"(?<!\\)\$", r"\\$", body)
        return fn(body, *args, **kwargs)
    return wrapper


for _name in ("markdown", "info", "success", "warning", "error", "caption", "write"):
    setattr(st, _name, _escape_dollars(getattr(st, _name)))
HERE = os.path.dirname(__file__)
EXAMPLE_FILE = os.path.join(HERE, "test_data", "BHCF20260930.txt")

STEPS = ["Home", "1 Upload", "2 Banks", "3 Combined", "4 Capital", "5 Approvals", "6 Cross-border", "7 Summary"]

# --------------------------------------------------------------------------------------
# Look and feel: large text, high contrast, big buttons (built for a busy senior lawyer)
# --------------------------------------------------------------------------------------
st.markdown("""
<style>
html, body, [class*="css"], .stMarkdown, p, li, label, .stRadio, .stSelectbox { font-size: 18px !important; }
h1 { font-size: 2.1rem !important; } h2 { font-size: 1.6rem !important; } h3 { font-size: 1.3rem !important; }
.block-container { padding-top: 3.6rem; max-width: 1250px; }
div.stButton > button { font-size: 16px !important; padding: 0.55rem 0.6rem; border-radius: 10px; width: 100%; }
.ai-box { border: 2px solid #6c4ee0; border-radius: 12px; padding: 14px 18px; background: rgba(108,78,224,0.06); margin: 8px 0 14px; }
.ai-badge { display: inline-block; background: #6c4ee0; color: white; font-size: 13px; font-weight: 700; padding: 2px 10px; border-radius: 99px; margin-bottom: 6px; }
.code-badge { display: inline-block; background: #0f5132; color: white; font-size: 13px; font-weight: 700; padding: 2px 10px; border-radius: 99px; }
.card { border: 1px solid rgba(128,128,128,0.35); border-radius: 14px; padding: 16px 18px; height: 100%; }
.card h4 { margin: 0 0 6px 0; font-size: 1.1rem; }
.st-key-tabbar div.stButton > button { padding: 0.45rem 0.2rem; }
.st-key-tabbar div.stButton > button p { font-size: 15px !important; white-space: nowrap; }
table.nice { width: 100%; border-collapse: collapse; font-size: 16px; margin: 6px 0 12px; }
table.nice th { text-align: left; background: #eef1f7; padding: 8px 10px; border-bottom: 2px solid #c9d0e0; }
table.nice td { padding: 8px 10px; border-bottom: 1px solid #e3e6ee; vertical-align: top; }
.big-cat { font-size: 2.4rem; font-weight: 800; line-height: 1.1; }
.muted { opacity: 0.75; font-size: 15px; }
</style>
""", unsafe_allow_html=True)


# --------------------------------------------------------------------------------------
# Shared state helpers
# --------------------------------------------------------------------------------------
@st.cache_resource
def load_library():
    return Library()


LIB = load_library()
S = st.session_state
S.setdefault("step", 0)
S.setdefault("saved", {})
S.setdefault("ai_cache", {})


def go_to(i):
    S["step"] = i


def saved(name, default=None):
    return S["saved"].get(name, default)


def keep(name, value):
    S["saved"][name] = value
    return value


def fingerprint(*parts):
    return hashlib.md5(json.dumps(parts, sort_keys=True, default=str).encode()).hexdigest()[:12]


def banks_ready():
    return bool(S.get("banks") and saved("acq") and saved("tgt"))


# --------------------------------------------------------------------------------------
# Sidebar: API credentials
#   >>> FOR SUBMISSION: replace this block with the required openai_credentials module
#       (https://github.com/stoneman-mls/openai_credentials). Anthropic is here only so you can
#       test cheaply with Claude Haiku while building.
# --------------------------------------------------------------------------------------
with st.sidebar:
    st.header("🔑 API credentials")
    provider = st.radio("AI provider", ["Anthropic (Claude)", "OpenAI"],
                        index=0 if saved("provider", "Anthropic (Claude)") == "Anthropic (Claude)" else 1)
    keep("provider", provider)
    env_key = os.environ.get("ANTHROPIC_API_KEY" if provider.startswith("Anthropic") else "OPENAI_API_KEY", "")
    key = st.text_input("API key", type="password", value=env_key,
                        help="Your key is held only in this browser session and never saved.")
    model = st.text_input("Model", value=ai.DEFAULT_MODELS[provider])
    S["provider"], S["api_key"], S["model"] = provider, key, model
    st.divider()
    counter_slot = st.empty()
    st.caption("Uploaded figures stay in this session. Only public filing data and regulation passages are sent "
               "to the AI provider. Do not upload confidential client information to a hosted version.")


# --------------------------------------------------------------------------------------
# Tab bar (highlights where you are; every tab is clickable) + Back / Next
# --------------------------------------------------------------------------------------
def tab_bar():
    cols = st.container(key="tabbar").columns(len(STEPS))
    for i, (col, label) in enumerate(zip(cols, STEPS)):
        col.button(label, key=f"tab_{i}", on_click=go_to, args=(i,),
                   type="primary" if S["step"] == i else "secondary", width="stretch")
    st.progress(S["step"] / (len(STEPS) - 1))


def back_next():
    st.divider()
    left, _, right = st.columns([1, 3, 1])
    if S["step"] > 0:
        left.button("⬅  Back", on_click=go_to, args=(S["step"] - 1,), width="stretch", key="back_btn")
    if S["step"] < len(STEPS) - 1:
        right.button("Next  ➡", on_click=go_to, args=(S["step"] + 1,), type="primary", width="stretch", key="next_btn")


def need_banks():
    st.info("👋 First, load the two banks in **Step 1 · Upload**.")
    st.button("Go to Step 1", on_click=go_to, args=(1,), type="primary")


# --------------------------------------------------------------------------------------
# Showing AI output: always badged, always with its sources one click away
# --------------------------------------------------------------------------------------
def show_ai(result, title="AI explanation"):
    if result is None:
        return
    check = result.get("check")
    note = ""
    if check is not None:
        note = (f" · cites {len(check['cited'])} passage(s), all verified as sent"
                if check["has_citations"] and not check["invalid"] else
                " · ⚠️ some citations were removed" if check["invalid"] else " · no citations given")
    st.markdown(f'<div class="ai-box"><span class="ai-badge">🤖 AI-GENERATED</span>'
                f'<span class="muted"> {title}{note}</span></div>', unsafe_allow_html=True)
    st.markdown(result["text"])
    if result.get("trail"):
        with st.expander("How this answer was checked (retrieve → answer → reflect → loop)"):
            for t in result["trail"]:
                st.write("• " + t)
    if result.get("passages"):
        with st.expander(f"📚 The {len(result['passages'])} regulation passages the AI was allowed to use"):
            for p in result["passages"]:
                st.markdown(f"**[{p['id']}] {p['cite']}** — {p['title']}  \n_{p['status']} · as of {p['as_of']}_")
                st.caption(p["text"][:900] + ("…" if len(p["text"]) > 900 else ""))


def run_ai(cache_key, fn, label="Asking the AI…"):
    """Run an AI step once per set of inputs; reuse the answer when the user comes back."""
    if cache_key in S["ai_cache"]:
        return S["ai_cache"][cache_key]
    if not ai.ready():
        st.warning("🔑 Add an API key in the sidebar to see the AI part of this step. All calculations above still work.")
        return None
    with st.spinner(label):
        try:
            out = fn()
        except Exception as exc:  # show a friendly message instead of a crash
            st.error(f"The AI call did not work: {exc}")
            return None
    S["ai_cache"][cache_key] = out
    return out


def question_box(step_name, facts):
    with st.expander("💬 Ask a question about this step (answers come only from the loaded regulations)"):
        q = st.text_input("Your question", key=f"q_{step_name}",
                          placeholder="e.g. What does the stress capital buffer apply to?")
        if st.button("Ask", key=f"ask_{step_name}") and q.strip():
            if not ai.ready():
                st.warning("Add an API key in the sidebar first.")
            else:
                with st.spinner("Searching the regulations, answering, then checking the answer…"):
                    S[f"ans_{step_name}"] = ai.ask(step_name, q, facts, LIB)
        if S.get(f"ans_{step_name}"):
            show_ai(S[f"ans_{step_name}"], "Answer")


def status_df(rows):
    df = pd.DataFrame(rows)
    return st.dataframe(df, hide_index=True, width="stretch")


# --------------------------------------------------------------------------------------
# Deal state used by steps 3-7 (works even if the user skips a step)
# --------------------------------------------------------------------------------------
def deal():
    banks = S["banks"]
    A, T = banks[saved("acq")], banks[saved("tgt")]
    a, t = A["latest"], T["latest"]
    ind = lambda who: {k: (saved(f"{who}_{k}") * 1e6 if saved(f"{who}_{k}") is not None else None)
                       for k in ("cross_jurisdictional", "nonbank_assets", "wstwf", "off_balance_sheet")}
    ia, it = ind("acq"), ind("tgt")
    ipf = {k: (ia[k] + it[k] if ia[k] is not None and it[k] is not None else None) for k in ia}
    thresholds = dict(E.CURRENT_THRESHOLDS)
    if saved("whatif_on"):
        thresholds.update(floor=saved("whatif_floor", 100) * 1e6, cat3_size=saved("whatif_cat3", 250) * 1e6)
    cat_a = E.determine_category(A["avg_assets"], ia, saved("acq_gsib", False), thresholds)
    cat_t = E.determine_category(T["avg_assets"], it, saved("tgt_gsib", False), thresholds)
    pf_avg = A["avg_assets"] + T["avg_assets"]
    cat_pf = E.determine_category(pf_avg, ipf, saved("acq_gsib", False), thresholds)
    tce = (t["equity"] or 0) - (t["preferred"] or 0) - (t["goodwill"] or 0)
    price = saved("price_bn", round(1.5 * tce / 1e6, 2)) * 1e6
    stock_pct = saved("stock_pct", 50)
    fv = saved("fv_bn", 0.0) * 1e6
    scb = saved("scb", 2.5)
    pf = E.pro_forma(a, t, price, stock_pct, fv)
    slr_exp = saved("slr_exposure_bn")
    tests = E.capital_tests(pf, cat_pf["category"], scb, slr_exp * 1e6 if slr_exp else None)
    return dict(A=A, T=T, a=a, t=t, cat_a=cat_a, cat_t=cat_t, cat_pf=cat_pf, pf_avg=pf_avg, ipf=ipf,
                thresholds=thresholds, tce=tce, price=price, stock_pct=stock_pct, fv=fv, scb=scb, pf=pf, tests=tests)


# ======================================================================================
# PAGES
# ======================================================================================
def page_home():
    st.title("🏦 Bank Merger Regulatory Check")
    st.markdown("#### Load two banks' public Federal Reserve filings. See which U.S. capital rules apply to each, "
                "what the deal switches on, how the price affects capital, and a summary you can send to a partner.")
    st.write("")
    cards = [
        ("1 · Upload", "Drop in the Fed's FR Y-9C data file and pick the acquirer and the target."),
        ("2 · Each bank", "Each bank's size category and the rules that apply to it, with citations."),
        ("3 · Combined", "The combined bank: which lines it crosses and what newly applies."),
        ("4 · Capital & price", "Move a slider to see how the purchase price changes capital ratios."),
        ("5 · Approvals", "Who approves the deal, and the 10% deposit and liabilities caps."),
        ("6 · Cross-border", "If a Canadian bank is buying: U.S. holding-company and OSFI capital checks."),
        ("7 · Summary", "One table of what is met, not met, newly applies, or can't be told — plus a draft memo."),
    ]
    rows = [cards[:4], cards[4:]]
    for row in rows:
        cols = st.columns(4)
        for col, (title, body) in zip(cols, row):
            col.markdown(f'<div class="card"><h4>{title}</h4>{body}</div>', unsafe_allow_html=True)
        st.write("")
    st.write("")
    c1, c2, _ = st.columns([1.2, 1.2, 2])
    c1.button("▶  Get started", type="primary", on_click=go_to, args=(1,), width="stretch")
    if c2.button("🧪  Try with example banks", width="stretch"):
        load_example()
        go_to(1)
        st.rerun()
    st.write("")
    st.markdown('<span class="code-badge">CODE</span> calculates and decides every number and every pass/fail. &nbsp; '
                '<span class="ai-badge">🤖 AI</span> reads the filings and regulations and explains — and only ever '
                'from the regulations loaded into this app (it cannot search the web).', unsafe_allow_html=True)
    with st.expander("📚 What regulations does the AI have? (its whole 'closed world')"):
        st.dataframe(pd.DataFrame(LIB.sources()), hide_index=True, width="stretch")
    st.caption("A screening tool, not legal advice. Figures are simplified; verify against the filings and rules cited.")


def load_example():
    with open(EXAMPLE_FILE, "rb") as fh:
        S["files"] = {"BHCF20260930.txt (example)": fh.read()}
    parse_files()
    names = {v["latest"]["name"]: k for k, v in S["banks"].items()}
    keep("acq", names.get("ASSOCIATED BANC-CORP"))
    keep("tgt", names.get("WAFD, INC."))


def parse_files():
    parsed = [E.parse_y9c_bulk(b) for b in S["files"].values()]
    S["banks"] = E.merge_quarters(parsed)
    S["quarters"] = [nice_date(q) for q in sorted({q for q, _ in parsed if q})]


def nice_date(q):
    import datetime
    try:
        return datetime.datetime.strptime(q, "%Y%m%d").strftime("%d %b %Y")
    except (TypeError, ValueError):
        return str(q)


def html_table(rows):
    """A plain table that wraps long text and uses large type (easier to read than a grid)."""
    df = pd.DataFrame(rows)
    st.markdown(df.to_html(index=False, classes="nice", border=0, escape=True), unsafe_allow_html=True)


def page_upload():
    st.header("Step 1 · Upload the banks' filings")
    st.markdown("Download the **FR Y-9C** data from the Federal Reserve's National Information Center "
                "([ffiec.gov/npw](https://www.ffiec.gov/npw/FinancialReport/FinancialDataDownload)) and upload it here. "
                "Upload **up to four quarters** — the size categories use a four-quarter average.")
    files = st.file_uploader("FR Y-9C bulk data files (.txt)", type=["txt"], accept_multiple_files=True)
    c1, _ = st.columns([1, 3])
    if c1.button("🧪 Use the example file instead"):
        load_example()
        st.rerun()
    if files:
        S["files"] = {f.name: f.getvalue() for f in files}
        parse_files()
    if not S.get("banks"):
        return
    banks = S["banks"]
    q = S.get("quarters", [])
    st.success(f"Read **{len(banks)}** holding companies from **{len(q)}** quarter(s): {', '.join(q)}.")
    if len(q) < 4:
        st.warning(f"Only {len(q)} quarter(s) loaded, so 'average assets' uses {len(q)} quarter(s) instead of four. "
                   "Results near a threshold may change once all four quarters are loaded.")

    order = sorted(banks, key=lambda k: -banks[k]["latest"]["total_assets"])
    label = lambda k: f"{banks[k]['latest']['name'].title()} — {banks[k]['latest']['city']}, {banks[k]['latest']['state']} — {E.bn(banks[k]['latest']['total_assets'])}"
    c1, c2 = st.columns(2)
    acq = c1.selectbox("🏦 Acquirer (the buyer)", order, format_func=label,
                       index=order.index(saved("acq")) if saved("acq") in order else 0)
    tgt = c2.selectbox("🎯 Target (being bought)", order, format_func=label,
                       index=order.index(saved("tgt")) if saved("tgt") in order else min(1, len(order) - 1))
    keep("acq", acq)
    keep("tgt", tgt)
    if acq == tgt:
        st.error("Pick two different banks.")
        return
    a, t = banks[acq]["latest"], banks[tgt]["latest"]

    st.subheader("What the app read from the filings")
    st.markdown('<span class="code-badge">CODE</span> read these figures directly from the Fed data (no AI involved). '
                'The last column shows where each one comes from.', unsafe_allow_html=True)
    rows = []
    for k, (codes, lbl, loc) in E.FIELDS.items():
        fmt = (lambda v: "n/a" if v is None else (f"{v:.2f}%" if k != "aoci_opt_out" else ("Yes" if v == 1 else "No"))) \
            if k in E.RATIO_FIELDS else E.bn
        rows.append({"Figure": lbl, "Acquirer": fmt(a[k]), "Target": fmt(t[k]), "Y-9C line (MDRM)": f"{' + '.join(codes)} · {loc}"})
    st.dataframe(pd.DataFrame(rows), hide_index=True, width="stretch", height=600)

    with st.expander("➕ Optional: FR Y-15 risk indicators (only matter if the combined bank is $100bn+)"):
        st.markdown("Banks with $100bn+ file the FR Y-15. These indicators can push a bank into a higher category even "
                    "below the size lines. Leave blank if unknown — the app will say what it couldn't check.")
        for who, nm in (("acq", a["name"].title()), ("tgt", t["name"].title())):
            st.markdown(f"**{nm}** (all in $ billions)")
            cols = st.columns(5)
            for col, (k, lbl) in zip(cols, [("wstwf", "Weighted short-term wholesale funding"), ("nonbank_assets", "Nonbank assets"),
                                            ("off_balance_sheet", "Off-balance-sheet exposure"), ("cross_jurisdictional", "Cross-jurisdictional activity")]):
                keep(f"{who}_{k}", col.number_input(lbl, min_value=0.0, value=saved(f"{who}_{k}"), key=f"w_{who}_{k}"))
            keep(f"{who}_gsib", cols[4].checkbox("U.S. GSIB", value=saved(f"{who}_gsib", False), key=f"w_{who}_gsib"))

    st.subheader("Data checks")
    c1, c2 = st.columns(2)
    for col, bank in ((c1, a), (c2, t)):
        with col:
            st.markdown(f"**{bank['name'].title()}**  <span class='code-badge'>CODE</span>", unsafe_allow_html=True)
            for chk in E.cross_checks(bank):
                icon = "✅" if chk["ok"] else ("❔" if chk["ok"] is None else "❌")
                st.write(f"{icon} {chk['check']} — {chk['result']}")

    st.subheader("AI review of the figures")
    c1, c2 = st.columns(2)
    for col, bank in ((c1, a), (c2, t)):
        with col:
            figs = "\n".join(f"- {E.FIELDS[k][1]}: {(f'{bank[k]:.2f}' if k in E.RATIO_FIELDS and bank[k] is not None else E.bn(bank[k], 2))}" for k in E.FIELDS)
            checks = "\n".join(f"- {c['check']}: {c['result']}" for c in E.cross_checks(bank))
            res = run_ai(fingerprint("review", bank["rssd"], bank["quarter"]),
                         lambda: {"text": ai.review_bank("1 Upload", bank["name"], figs, checks)},
                         f"AI is reviewing {bank['name'].title()}…")
            st.markdown(f"**{bank['name'].title()}**")
            show_ai(res, "Data review (figures only, no law)")


def page_each_bank():
    st.header("Step 2 · Each bank on its own")
    if not banks_ready():
        return need_banks()
    d = deal()
    c1, c2 = st.columns(2)
    for col, who, B, cat in ((c1, "Acquirer", d["A"], d["cat_a"]), (c2, "Target", d["T"], d["cat_t"])):
        with col:
            b = B["latest"]
            st.markdown(f"#### {who}: {b['name'].title()}")
            st.markdown(f"<div class='big-cat'>Category {cat['category'] if cat['category'] != 'None' else '— none (under $100bn)'}</div>",
                        unsafe_allow_html=True)
            st.markdown(f"<span class='code-badge'>CODE</span> {cat['reason']}  \n<span class='muted'>Source: {cat['cite']} · "
                        f"average of {B['quarters']} quarter(s): {E.bn(B['avg_assets'])}</span>", unsafe_allow_html=True)
            if cat["missing"]:
                st.caption("Not provided: " + ", ".join(cat["missing"]) + " (add them in Step 1 if available).")
            html_table(E.requirements_for(cat["category"]))
    facts = (f"Acquirer {d['a']['name']}: avg assets {E.bn(d['A']['avg_assets'])}, category {d['cat_a']['category']} ({d['cat_a']['reason']})\n"
             f"Target {d['t']['name']}: avg assets {E.bn(d['T']['avg_assets'])}, category {d['cat_t']['category']} ({d['cat_t']['reason']})")
    st.subheader("What this means")
    res = run_ai(fingerprint("each", facts),
                 lambda: ai.explain("2 Each bank",
                                    "Explain in under 180 words why each bank falls in its category and the most important "
                                    "rules that follow for it. Use the category definitions in the passages.",
                                    facts, "categorization banking organizations Category IV average total consolidated assets "
                                    "$100 billion $250 billion cross-jurisdictional activity weighted short-term wholesale funding", LIB))
    show_ai(res)
    question_box("2 Each bank", facts)


def page_combined():
    st.header("Step 3 · The combined bank")
    if not banks_ready():
        return need_banks()
    d = deal()
    before, after = d["cat_a"]["category"], d["cat_pf"]["category"]
    c1, c2, c3 = st.columns(3)
    c1.metric("Acquirer today", f"Cat. {before}" if before != "None" else "Under $100bn", E.bn(d["A"]["avg_assets"]), delta_color="off")
    c2.metric("Combined bank", f"Cat. {after}" if after != "None" else "Under $100bn", E.bn(d["pf_avg"]), delta_color="off")
    line, gap = E.next_line(d["pf_avg"], d["thresholds"])
    c3.metric("Room before the next line", E.bn(gap) if gap else "—", f"next: {line}" if line else "", delta_color="off")
    st.markdown(f"<span class='code-badge'>CODE</span> {d['cat_pf']['reason']}", unsafe_allow_html=True)
    if CATEGORY_RANK(after) > CATEGORY_RANK(before):
        st.error(f"⚠️ The deal moves the acquirer from **{label_cat(before)}** to **{label_cat(after)}**.")
    else:
        st.success("The deal does not move the acquirer into a new category.")
    if d["cat_pf"]["missing"] and after != "None":
        st.warning("Some FR Y-15 risk indicators were not provided, so a higher category cannot be ruled out.")
    st.caption("Size is the sum of each bank's average assets. In practice the four-quarter average catches up over the "
               "quarters after closing, but supervisors expect the new category's infrastructure to be ready at closing.")

    st.subheader("What the deal switches on")
    rows = E.newly_applying(before, after)
    if rows:
        html_table(rows)
    else:
        st.write("Nothing new — the same requirements apply before and after.")

    with st.expander("🔧 What-if: move the size lines (e.g. if thresholds are indexed to GDP)"):
        st.markdown("The **$100bn** and **$250bn** lines are currently in force. Proposals to index them have been discussed; "
                    "**no indexing proposal is in this app's library**, so these sliders are a hypothetical stress test only.")
        on = st.toggle("Use my what-if lines", value=saved("whatif_on", False), key="w_whatif_on")
        keep("whatif_on", on)
        keep("whatif_floor", st.slider("Category IV line ($bn)", 100, 200, saved("whatif_floor", 100), 5, key="w_wf"))
        keep("whatif_cat3", st.slider("Category III line ($bn)", 250, 500, saved("whatif_cat3", 250), 10, key="w_wc"))

    facts = (f"Acquirer {d['a']['name']} category {before}; target {d['t']['name']} category {d['cat_t']['category']}; "
             f"combined average assets {E.bn(d['pf_avg'])} -> category {after}. Reason: {d['cat_pf']['reason']}\n"
             f"Requirements that change: {json.dumps(rows) if rows else 'none'}\n"
             f"Next threshold: {line} with {E.bn(gap) if gap else 'n/a'} of room.")
    st.subheader("What this means for the deal")
    res = run_ai(fingerprint("combined", facts),
                 lambda: ai.explain("3 Combined",
                                    "In under 200 words, explain to the deal lawyer what the combination does to the acquirer's "
                                    "regulatory position: the category result, what newly applies (if anything), and how close the "
                                    "combined bank is to the next threshold. Cite the passages.",
                                    facts, "Category IV Category III $100 billion $250 billion stress capital buffer supplementary "
                                    "leverage ratio capital plan applicability", LIB))
    show_ai(res)
    question_box("3 Combined", facts)


def fmt_pct(v):
    return "n/a" if v is None else f"{v:.2f}%"


def CATEGORY_RANK(c):
    return E.CATEGORY_ORDER.index(c)


def label_cat(c):
    return "below $100bn" if c == "None" else f"Category {c}"


def page_capital():
    st.header("Step 4 · Capital after closing — and what the price does")
    if not banks_ready():
        return need_banks()
    d0 = deal()
    tce = d0["tce"]
    st.markdown(f"The target's **tangible common equity** is {E.bn(tce, 2)}. Anything paid above it (after any fair-value "
                "write-down) becomes **goodwill**, which is **deducted from CET1** (12 CFR 217.22(a)(1)) while the target's "
                "whole balance sheet is added to the denominators.")
    c1, c2 = st.columns([2, 1])
    with c1:
        lo, hi = round(0.25 * tce / 1e6, 2), round(3.5 * tce / 1e6, 2)
        price = st.slider("💲 Purchase price for the target ($bn)", lo, hi, min(max(saved("price_bn", round(1.5 * tce / 1e6, 2)), lo), hi),
                          0.05, key="w_price")
        keep("price_bn", price)
        st.caption(f"= {price * 1e6 / tce:.2f}× the target's tangible common equity")
    with c2:
        keep("stock_pct", st.slider("Paid in acquirer shares (%)", 0, 100, saved("stock_pct", 50), 5, key="w_stock"))
    c1, c2, c3 = st.columns(3)
    keep("fv_bn", c1.number_input("Fair-value write-down of target assets at closing ($bn)", 0.0, value=saved("fv_bn", 0.0), step=0.1,
                                  key="w_fv", help="e.g. unrealised losses on the target's securities that purchase accounting crystallises"))
    d = deal()
    cat = d["cat_pf"]["category"]
    if cat != "None":
        keep("scb", c2.number_input("Stress capital buffer (%)", 2.5, 10.0, saved("scb", 2.5), 0.1, key="w_scb",
                                    help="Minimum 2.5% (12 CFR 225.8). Use the acquirer's published SCB."))
    if cat in ("III", "II", "I"):
        keep("slr_exposure_bn", c3.number_input("Total leverage exposure for SLR ($bn, optional)", 0.0,
                                                value=saved("slr_exposure_bn"), key="w_slr"))
    d = deal()
    pf, a = d["pf"], d["a"]

    st.subheader("Ratios before and after")
    cols = st.columns(4)
    for col, (k, nm) in zip(cols, [("cet1_ratio", "CET1"), ("tier1_ratio", "Tier 1"), ("total_ratio", "Total capital"), ("leverage_ratio", "Leverage")]):
        col.metric(f"{nm} ratio", f"{pf[k]:.2f}%", f"{pf[k] - a[k]:+.2f} pts vs acquirer today")
    st.dataframe(pd.DataFrame([{
        "Test": r["test"], "After the deal": "n/a" if r["value"] is None else f"{r['value']:.2f}%",
        "Minimum": f"{r['minimum']:.1f}%", "Minimum + buffer": f"{r['with_buffer']:.1f}%",
        "Result": E.STATUS_LABEL[r["status"]], "Source": r["cite"]} for r in d["tests"]]), hide_index=True, width="stretch")
    st.caption(f"New goodwill {E.bn(pf['new_goodwill'], 2)} · new shares {E.bn(pf['stock'], 2)} · cash paid {E.bn(pf['cash'], 2)} · "
               f"combined RWA {E.bn(pf['rwa'])}. Buffer used: {E.capital_requirements(cat, d['scb'])['buffer_name']}.")

    xs, ys, need, breakeven = E.price_sweep(a, d["t"], d["stock_pct"], d["fv"], cat, d["scb"])
    fig = go.Figure()
    fig.add_trace(go.Scatter(x=[x / 1e6 for x in xs], y=ys, mode="lines", name="CET1 ratio after the deal", line=dict(width=4)))
    fig.add_hline(y=need, line_dash="dash", annotation_text=f"Minimum + buffer {need:.1f}%", annotation_position="bottom left")
    fig.add_vline(x=d["price"] / 1e6, line_dash="dot", annotation_text="Your price")
    fig.update_layout(height=380, margin=dict(l=10, r=10, t=30, b=10), xaxis_title="Purchase price ($bn)",
                      yaxis_title="CET1 ratio (%)", font=dict(size=16), showlegend=False)
    st.plotly_chart(fig, width="stretch")
    if breakeven:
        st.info(f"📏 Above about **{E.bn(breakeven, 2)}** ({breakeven / tce:.2f}× tangible book), the combined CET1 ratio falls "
                f"inside the {need:.1f}% minimum-plus-buffer, which would restrict dividends and bonuses (12 CFR 217.11).")
    else:
        st.success("Across the whole price range shown, CET1 stays above the minimum plus buffer.")
    with st.expander("How this is calculated (simplified — check before relying on it)"):
        st.markdown(E.pro_forma.__doc__)

    facts = (f"Price {E.bn(d['price'], 2)} ({d['price'] / tce:.2f}x target TCE), {d['stock_pct']}% in stock. New goodwill "
             f"{E.bn(pf['new_goodwill'], 2)}. Combined category {cat}. Results: " +
             "; ".join(f"{r['test']} {fmt_pct(r['value'])} vs {r['with_buffer']:.1f}% -> {r['status']}" for r in d["tests"]) +
             f". Break-even price {E.bn(breakeven, 2) if breakeven else 'none in range'}.")
    st.subheader("What this means")
    key = fingerprint("capital", facts)
    if key in S["ai_cache"] or st.button("🤖 Explain these results", type="primary"):
        res = run_ai(key, lambda: ai.explain(
            "4 Capital", "In under 170 words explain what the price does to the combined bank's capital: the goodwill "
            "deduction, which ratio is tightest, and what falling inside the buffer would mean. Cite the passages.",
            facts, "goodwill deduction common equity tier 1 capital conservation buffer maximum payout ratio stress capital buffer", LIB))
        show_ai(res)
    else:
        st.caption("The AI explanation runs when you press the button, so moving the slider stays fast and free.")
    question_box("4 Capital", facts)


def page_approvals():
    st.header("Step 5 · Approvals overview and statutory caps")
    if not banks_ready():
        return need_banks()
    d = deal()
    c1, c2 = st.columns(2)
    structure = c1.radio("Deal structure", ["Holding company buys holding company", "Bank merges into bank", "Both"],
                         index=["Holding company buys holding company", "Bank merges into bank", "Both"].index(saved("structure", "Both")), key="w_structure")
    keep("structure", structure)
    charters = ["National bank (OCC)", "State member bank (Fed)", "State non-member bank (FDIC)"]
    charter = c2.selectbox("Charter of the resulting bank", charters, index=charters.index(saved("charter", charters[0])), key="w_charter")
    keep("charter", charter)
    st.subheader("Who approves")
    html_table(E.approval_path(structure, charter))

    st.subheader("The two hard 10% caps")
    c1, c2 = st.columns(2)
    keep("us_deposits_bn", c1.number_input("Total deposits of all U.S. insured institutions ($bn)", 0.0, value=saved("us_deposits_bn"),
                                           key="w_usdep", help="Take this from FDIC data for the latest quarter."))
    keep("fin_liab_bn", c2.number_input("Aggregate financial-sector liabilities published by the Fed ($bn)", 0.0,
                                        value=saved("fin_liab_bn"), key="w_finliab", help="Published annually by July 1 (12 CFR 251)."))
    dep = E.deposit_cap(d["pf"]["deposits"], (saved("us_deposits_bn") or 0) * 1e6)
    liab = E.liabilities_cap(d["pf"]["total_liabilities"], (saved("fin_liab_bn") or 0) * 1e6)
    c1.markdown(f"**Nationwide deposit cap** — {E.STATUS_LABEL[dep['status']]}  \nCombined deposits {E.bn(d['pf']['deposits'])}. {dep['detail']}  \n"
                "<span class='muted'>12 U.S.C. 1842(d)(2)(A); 12 U.S.C. 1828(c)(13)</span>", unsafe_allow_html=True)
    c2.markdown(f"**Liabilities concentration limit** — {E.STATUS_LABEL[liab['status']]}  \nCombined liabilities {E.bn(d['pf']['total_liabilities'])}. "
                f"{liab['detail']}  \n<span class='muted'>12 U.S.C. 1852(b); 12 CFR 251</span>", unsafe_allow_html=True)
    S["caps"] = {"deposit": dep, "liabilities": liab}

    facts = (f"Structure: {structure}; resulting bank charter: {charter}. Approval path: {json.dumps(E.approval_path(structure, charter))}. "
             f"Deposit cap: {dep['status']} ({dep['detail']}). Liabilities cap: {liab['status']} ({liab['detail']}).")
    st.subheader("What the approving agency will look at")
    res = run_ai(fingerprint("approvals", facts),
                 lambda: ai.explain("5 Approvals",
                                    "In under 180 words, list the statutory factors the approving agency must consider for this "
                                    "structure (competition, financial and managerial resources, convenience and needs, financial "
                                    "stability, anti-money-laundering effectiveness, as stated in the passages). Brief overview only.",
                                    facts, "factors consider application competitive effects convenience needs community financial "
                                    "managerial resources stability United States banking money laundering", LIB))
    show_ai(res)
    question_box("5 Approvals", facts)


def page_cross_border():
    st.header("Step 6 · Cross-border: a Canadian bank buying a U.S. bank")
    if not banks_ready():
        return need_banks()
    d = deal()
    on = st.toggle("The acquirer is owned by a Canadian bank (the uploaded acquirer is its U.S. holding company)",
                   value=saved("xb_on", False), key="w_xb_on")
    keep("xb_on", on)
    if not on:
        st.info("Not a cross-border deal — nothing to check here. Press **Next** to go to the summary.")
        return
    st.subheader("The Canadian parent's figures")
    st.markdown("Upload the parent's **Supplementary Regulatory Capital Disclosure** or **Pillar 3** PDF and the AI will pull the "
                "figures out (you can correct them), or type them in. All amounts in **C$ millions**.")
    pdf = st.file_uploader("Parent capital disclosure (PDF)", type=["pdf"], key="w_parent_pdf")
    if pdf is not None and st.button("🤖 Read the figures from this PDF", type="primary"):
        from pypdf import PdfReader
        text = ai.find_capital_pages(PdfReader(io.BytesIO(pdf.getvalue())))
        out = run_ai(fingerprint("extract", pdf.name, len(text)), lambda: {"data": ai.extract_parent_figures("6 Cross-border", text)},
                     "AI is reading the PDF…")
        if out:
            figs = out["data"].get("figures", {})
            for k in ai.PARENT_FIELDS:
                v = (figs.get(k) or {}).get("value")
                if v is not None:
                    keep(f"p_{k}", float(v))
            S["extract_quotes"] = {k: (figs.get(k) or {}).get("found_as") for k in ai.PARENT_FIELDS}
            st.rerun()
    if S.get("extract_quotes"):
        with st.expander("Where the AI found each figure (check these)"):
            for k, qt in S["extract_quotes"].items():
                st.write(f"• **{ai.PARENT_FIELDS[k]}**: {qt or 'not found'}")
    c = st.columns(5)
    for col, (k, lbl, default) in zip(c, [("total_assets", "Total assets", 2000000.0), ("cet1", "CET1 capital", 95000.0),
                                          ("tier1", "Tier 1 capital", 110000.0), ("rwa", "Risk-weighted assets", 720000.0),
                                          ("leverage_exposure", "Leverage exposure", 2400000.0)]):
        keep(f"p_{k}", col.number_input(lbl, 0.0, value=float(saved(f"p_{k}", default)), step=1000.0, key=f"w_p_{k}"))
    if not S.get("extract_quotes"):
        st.caption("⚠️ The starting numbers are placeholders, not a real bank. Replace them with the parent's figures.")
    c = st.columns(4)
    keep("fx", c[0].number_input("C$ per US$1", 0.5, 3.0, saved("fx", 1.40), 0.01, key="w_fx", help="Enter the current exchange rate."))
    keep("dsib", c[1].checkbox("Parent is a D-SIB", value=saved("dsib", True), key="w_dsib"))
    keep("dsb", c[2].number_input("Domestic Stability Buffer (%)", 0.0, 4.0, saved("dsb", 3.0), 0.25, key="w_dsb",
                                  help="Set by OSFI and not in this app's library — check OSFI's latest announcement."))
    keep("branch_bn", c[3].number_input("Parent's U.S. branch assets (US$bn)", 0.0, value=saved("branch_bn", 0.0), key="w_branch"))

    fx = saved("fx")
    us_nonbranch = d["a"]["total_assets"] + d["t"]["total_assets"]
    global_us = saved("p_total_assets") * 1e3 / fx  # C$m -> US$ thousands
    ihc = E.ihc_test(global_us, us_nonbranch, before_nonbranch=d["a"]["total_assets"])
    parent = {k: saved(f"p_{k}") for k in ("cet1", "tier1", "rwa", "leverage_exposure")}
    to_cad = lambda usd_thousands: usd_thousands / 1e3 * fx
    pp = E.parent_pro_forma(parent, to_cad(d["pf"]["new_goodwill"]), to_cad(d["t"]["rwa"]), to_cad(d["t"]["total_assets"]),
                            to_cad(d["pf"]["stock"]), to_cad(d["pf"]["cash"]), saved("dsb"), saved("dsib"))

    st.subheader("Two regimes at once")
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("#### 🇺🇸 U.S. intermediate holding company")
        st.markdown(f"**IHC requirement** — {E.STATUS_LABEL[ihc['status']]}  \n{ihc['detail']}  \n<span class='muted'>12 CFR 252.153</span>",
                    unsafe_allow_html=True)
        st.markdown(f"**U.S. category of the IHC after the deal:** {label_cat(d['cat_pf']['category'])} (from Step 3).  \n"
                    f"Combined U.S. operations incl. branches: {E.bn(us_nonbranch + saved('branch_bn') * 1e6)} — branch assets count here "
                    "but not toward the $50bn IHC trigger.")
    with c2:
        st.markdown("#### 🇨🇦 Parent under OSFI")
        rows = [("CET1 ratio", pp["cet1_ratio_before"], pp["cet1_ratio_after"], pp["cet1_target"], "OSFI CAR (2027) Ch.1 s.10"),
                ("Tier 1 ratio", pp["tier1_ratio_before"], pp["tier1_ratio_after"], pp["tier1_target"], "OSFI CAR (2027) Ch.1 s.10"),
                ("Leverage ratio", pp["leverage_before"], pp["leverage_after"], pp["leverage_target"], "OSFI LR Guideline")]
        S["parent_rows"] = [{"test": r[0], "after": r[2], "target": r[3], "status": "met" if r[2] >= r[3] else "not_met", "cite": r[4]} for r in rows]
        st.dataframe(pd.DataFrame([{"Ratio": r[0], "Before": f"{r[1]:.2f}%", "After": f"{r[2]:.2f}%", "OSFI target": f"{r[3]:.2f}%",
                                    "Result": E.STATUS_LABEL["met" if r[2] >= r[3] else "not_met"], "Source": r[4]} for r in rows]),
                     hide_index=True, width="stretch")
        st.caption("Simplified: the target's U.S. RWA is used as-is; OSFI risk weights may differ. Goodwill is deducted from CET1.")
    S["ihc"] = ihc

    facts = (f"Parent (Canadian, D-SIB={saved('dsib')}) CET1 ratio {pp['cet1_ratio_before']:.2f}% -> {pp['cet1_ratio_after']:.2f}% vs target "
             f"{pp['cet1_target']:.2f}%; leverage {pp['leverage_before']:.2f}% -> {pp['leverage_after']:.2f}% vs {pp['leverage_target']:.2f}%. "
             f"U.S. non-branch assets after deal {E.bn(us_nonbranch)}; IHC: {ihc['detail']} U.S. category {d['cat_pf']['category']}.")
    st.subheader("What this means")
    res = run_ai(fingerprint("xb", facts),
                 lambda: ai.explain("6 Cross-border",
                                    "In under 200 words, explain the dual-compliance position: whether the U.S. intermediate holding "
                                    "company requirement applies and what that brings, and how the deal affects the parent's OSFI ratios "
                                    "against its targets. Cite U.S. and OSFI passages.",
                                    facts, "U.S. intermediate holding company foreign banking organization $50 billion non-branch assets "
                                    "D-SIB target capital ratios leverage ratio buffer", LIB, k=8))
    show_ai(res)
    question_box("6 Cross-border", facts)


def summary_rows(d):
    rows = []
    before, after = d["cat_a"]["category"], d["cat_pf"]["category"]
    for r in E.newly_applying(before, after):
        rows.append({"Status": E.STATUS_LABEL["newly_applies"], "Check": r["Requirement"],
                     "Detail": f"{r['Before (acquirer alone)']} → {r['After the deal']}", "Source": r["Source"]})
    rows.append({"Status": E.STATUS_LABEL["newly_applies"] if CATEGORY_RANK(after) > CATEGORY_RANK(before) else E.STATUS_LABEL["met"],
                 "Check": "Size category after the deal", "Detail": f"{label_cat(before)} → {label_cat(after)} ({E.bn(d['pf_avg'])})",
                 "Source": d["cat_pf"]["cite"]})
    if d["cat_pf"]["missing"] and after != "None":
        rows.append({"Status": E.STATUS_LABEL["unknown"], "Check": "FR Y-15 risk indicators",
                     "Detail": "Not provided: " + ", ".join(d["cat_pf"]["missing"]), "Source": "12 CFR 252.5"})
    for r in d["tests"]:
        rows.append({"Status": E.STATUS_LABEL[r["status"]], "Check": f"{r['test']} after the deal",
                     "Detail": "n/a" if r["value"] is None else f"{r['value']:.2f}% (min {r['minimum']:.1f}%, with buffer {r['with_buffer']:.1f}%)",
                     "Source": r["cite"]})
    caps = S.get("caps") or {"deposit": E.deposit_cap(d["pf"]["deposits"], 0), "liabilities": E.liabilities_cap(d["pf"]["total_liabilities"], 0)}
    rows.append({"Status": E.STATUS_LABEL[caps["deposit"]["status"]], "Check": "10% nationwide deposit cap", "Detail": caps["deposit"]["detail"],
                 "Source": "12 U.S.C. 1842(d)(2)(A); 1828(c)(13)"})
    rows.append({"Status": E.STATUS_LABEL[caps["liabilities"]["status"]], "Check": "10% liabilities concentration limit",
                 "Detail": caps["liabilities"]["detail"], "Source": "12 U.S.C. 1852(b); 12 CFR 251"})
    if saved("xb_on") and S.get("ihc"):
        rows.append({"Status": E.STATUS_LABEL[S["ihc"]["status"]], "Check": "U.S. intermediate holding company requirement",
                     "Detail": S["ihc"]["detail"], "Source": "12 CFR 252.153"})
        for r in S.get("parent_rows", []):
            rows.append({"Status": E.STATUS_LABEL[r["status"]], "Check": f"Parent {r['test']} (OSFI)",
                         "Detail": f"{r['after']:.2f}% vs target {r['target']:.2f}%", "Source": r["cite"]})
    q = len(S.get("quarters", []))
    if q < 4:
        rows.append({"Status": E.STATUS_LABEL["unknown"], "Check": "Four-quarter averages",
                     "Detail": f"Only {q} quarter(s) loaded", "Source": "12 CFR 252.5"})
    return rows


def page_summary():
    st.header("Step 7 · Summary")
    if not banks_ready():
        return need_banks()
    d = deal()
    st.markdown(f"**{d['a']['name'].title()}** acquiring **{d['t']['name'].title()}** · price {E.bn(d['price'], 2)} "
                f"({d['stock_pct']}% stock) · data as of {', '.join(S.get('quarters', []))}")
    rows = summary_rows(d)
    counts = pd.Series([r["Status"] for r in rows]).value_counts()
    cols = st.columns(4)
    for col, s in zip(cols, ["met", "not_met", "newly_applies", "unknown"]):
        lbl = E.STATUS_LABEL[s]
        n = int(counts.get(lbl, 0)) + (int(counts.get(E.STATUS_LABEL["warning"], 0)) if s == "newly_applies" else 0)
        col.metric(lbl, n)
    st.markdown("<span class='code-badge'>CODE</span> Every status below was decided by the app's rules, not by the AI.", unsafe_allow_html=True)
    html_table(rows)

    facts = "\n".join(f"- {r['Status']} | {r['Check']} | {r['Detail']} | {r['Source']}" for r in rows)
    facts = f"Deal: {d['a']['name']} acquiring {d['t']['name']}, price {E.bn(d['price'], 2)}.\nSUMMARY TABLE:\n{facts}"
    st.subheader("Draft memo for the deal team")
    key = fingerprint("memo", facts)

    def memo_and_check():
        draft = ai.explain("7 Summary",
                           "Draft a short internal memo (under 300 words) to a banking M&A partner: headline conclusion, the items "
                           "that newly apply or are not met, the items that could not be checked and what data is needed, and "
                           "next steps. Use the SUMMARY TABLE statuses exactly; cite passages for each legal point.",
                           facts, "Category III Category IV stress capital buffer supplementary leverage ratio capital plan "
                           "deposits 10 percent concentration limit goodwill deduction", LIB, k=8)
        draft["verdict"] = ai.reflect("7 Summary", draft["text"], draft["passages"], "Is the memo supported by its sources?")
        return draft

    res = run_ai(key, memo_and_check, "AI is drafting the memo, then a second AI call is checking it…")
    if res:
        show_ai(res, "Draft memo")
        v = res.get("verdict", {})
        if v.get("supported"):
            st.success("🔎 Reflection check (second AI call): every legal statement is supported by its cited passages.")
        else:
            st.warning("🔎 Reflection check flagged these sentences as not clearly supported:\n\n" +
                       "\n".join(f"- {s}" for s in v.get("unsupported", [])) or "(no detail given)")
    c1, c2 = st.columns(2)
    c1.download_button("⬇️ Download summary + memo (Word)", make_docx(d, rows, res), file_name="merger_regulatory_summary.docx",
                       type="primary", width="stretch")
    c2.download_button("⬇️ Download AI audit trail", json.dumps(S.get("audit", []), indent=2),
                       file_name="ai_audit_trail.json", width="stretch")
    question_box("7 Summary", facts)


def make_docx(d, rows, memo):
    from docx import Document
    doc = Document()
    doc.add_heading("Bank merger regulatory screen", 0)
    doc.add_paragraph(f"{d['a']['name'].title()} acquiring {d['t']['name'].title()} — purchase price {E.bn(d['price'], 2)}, "
                      f"{d['stock_pct']}% in stock. Filing data as of {', '.join(S.get('quarters', []))}.")
    doc.add_heading("Results (decided by the app's rules)", 1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(["Status", "Check", "Detail", "Source"]):
        table.rows[0].cells[i].text = h
    for r in rows:
        cells = table.add_row().cells
        for i, k in enumerate(["Status", "Check", "Detail", "Source"]):
            cells[i].text = str(r[k])
    if memo:
        doc.add_heading("Draft memo (AI-generated — verify before use)", 1)
        for para in memo["text"].split("\n"):
            if para.strip():
                doc.add_paragraph(para.strip().lstrip("#").strip())
        doc.add_heading("Sources the AI was given", 2)
        for p in memo["passages"]:
            doc.add_paragraph(f"[{p['id']}] {p['cite']} — {p['source']} ({p['status']}, as of {p['as_of']})", style="List Bullet")
    doc.add_paragraph("Screening tool output, not legal advice. Calculations are simplified; verify against the filings and the rules cited.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ======================================================================================
tab_bar()
[page_home, page_upload, page_each_bank, page_combined, page_capital, page_approvals, page_cross_border, page_summary][S["step"]]()
back_next()
counter_slot.metric("AI calls this session", S.get("ai_calls", 0),
                    help="The assignment requires at least 5 AI calls per run. A full run through all steps makes 7 or more.")
